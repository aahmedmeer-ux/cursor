#!/usr/bin/env python3
"""Generate Usama Mir Ph.D. application Word package (cover letter + resume)."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

OUT_DIR = Path(__file__).resolve().parent
PHOTO = OUT_DIR / "usama_mir_headshot.png"
OUTPUT = OUT_DIR / "Usama_Mir_Oakland_University_PhD_Application.docx"
ARTIFACT = Path("/opt/cursor/artifacts/Usama_Mir_Oakland_University_PhD_Application.docx")

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x8B, 0x1E, 0x2D)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x44, 0x44, 0x44)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=6, line=1.08, space_after_exact=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after if space_after_exact is None else space_after_exact)
    pf.line_spacing = line


def add_horizontal_line(paragraph, color="1B2A4A", thickness="12"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), thickness)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=4, line=1.0)
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=NAVY)
    add_horizontal_line(p, color="8B1E2D", thickness="10")
    return p


def body_para(doc, text, size=10.5, after=4, before=0, justify=True):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=1.08)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=size, color=BODY)
    return p


def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=2, line=1.05)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=size, color=BODY)
    return p


def role_header(doc, title, org, dates):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=0, line=1.0)
    r1 = p.add_run(title)
    set_run_font(r1, size=10.5, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=2, line=1.0)
    # tab stop for dates on right
    tab_stops = p2.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r2 = p2.add_run(org)
    set_run_font(r2, size=10, italic=True, color=MUTED)
    r3 = p2.add_run("\t" + dates)
    set_run_font(r3, size=9.5, color=MUTED)
    return p


def edu_block(doc, degree, school, dates, extra=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=5, after=0, line=1.0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(degree)
    set_run_font(r1, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run("\t" + dates)
    set_run_font(r2, size=9.5, color=MUTED)
    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=1, line=1.0)
    r3 = p2.add_run(school)
    set_run_font(r3, size=10, italic=True, color=MUTED)
    if extra:
        body_para(doc, extra, size=10, after=2, before=1)


def configure_page(section):
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)


def add_header_block(doc, include_photo=True):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    table.columns[0].width = Inches(5.85)
    table.columns[1].width = Inches(1.35)

    left = table.cell(0, 0)
    right = table.cell(0, 1)

    # Clear default paragraphs
    left.paragraphs[0].clear()
    name_p = left.paragraphs[0]
    set_paragraph_spacing(name_p, before=0, after=2, line=1.0)
    name_run = name_p.add_run("USAMA MIR")
    set_run_font(name_run, name="Calibri", size=22, bold=True, color=NAVY)

    title_p = left.add_paragraph()
    set_paragraph_spacing(title_p, before=0, after=4, line=1.0)
    title_run = title_p.add_run(
        "Electrical & Telecommunications Engineer | RF Systems · Applied EM · Scientific ML"
    )
    set_run_font(title_run, size=9.5, italic=True, color=ACCENT)

    contact_p = left.add_paragraph()
    set_paragraph_spacing(contact_p, before=0, after=0, line=1.05)
    contact = (
        "Islamabad, Pakistan  |  payoneerusama@gmail.com  |  "
        "linkedin.com/in/usama-mir-08678468"
    )
    c_run = contact_p.add_run(contact)
    set_run_font(c_run, size=9, color=MUTED)

    target_p = left.add_paragraph()
    set_paragraph_spacing(target_p, before=2, after=0, line=1.0)
    t_run = target_p.add_run(
        "Ph.D. Applicant — Electrical & Computer Engineering, Oakland University"
    )
    set_run_font(t_run, size=9, bold=True, color=NAVY)

    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if include_photo and PHOTO.exists():
        run = right.paragraphs[0].add_run()
        run.add_picture(str(PHOTO), width=Inches(1.22))

    # thin line under header
    line_p = doc.add_paragraph()
    set_paragraph_spacing(line_p, before=4, after=2, line=1.0)
    add_horizontal_line(line_p, color="1B2A4A", thickness="18")
    return table


def build_resume(doc):
    add_header_block(doc, include_photo=True)

    section_heading(doc, "Academic & Research Summary")
    body_para(
        doc,
        "Electrical and Telecommunications Engineer specializing in RF systems, applied "
        "electromagnetics, counter-UAS sensing, and scientific machine learning. Demonstrated "
        "expertise in X-band sensor development, satellite array antenna design, anti-drone "
        "operator detection, and 6G-oriented network architectures. Proven ability to integrate "
        "Physics-Informed Neural Networks (PINNs), spatial digital twins (GSLAM), and Siamese "
        "CNNs for RF fingerprinting, multipath-resilient emitter identification, and fluid-"
        "dynamics modeling. Seeking Ph.D. study at Oakland University to advance research in "
        "applied electromagnetics, computational sensing, and intelligent RF systems aligned "
        "with the Applied EMAG and Wireless Laboratory and related SECS faculty strengths.",
        size=10,
        after=2,
    )

    section_heading(doc, "Education")
    edu_block(
        doc,
        "Master of Science in Electrical Engineering (MSEE)",
        "National University of Sciences and Technology (NUST), Islamabad, Pakistan",
        "2025 – 2026",
        "Thesis: Hybrid GSLAM, PINNs & Siamese CNN Framework for Complex Real-World Radio "
        "Frequency Fingerprinting. Focus: differentiable ray tracing (NVIDIA Sionna RT), "
        "digital-twin CIR modeling, multipath suppression, and hardware-signature extraction "
        "for 6G-related rogue-emitter identification and power allocation.",
    )
    edu_block(
        doc,
        "Bachelor of Science in Telecommunications Engineering",
        "University of Engineering and Technology (UET), Peshawar, Pakistan",
        "2013 – 2017",
        "Foundational training in electromagnetics, wireless communications, signal processing, "
        "and RF systems. Early research exposure to WLAN optimization, smart antennas, and "
        "high-frequency antenna concepts.",
    )

    section_heading(doc, "Publications & Research Under Review")
    body_para(
        doc,
        "Mir, U., et al. “A Hybrid TSA-DG Physics-Informed Neural Network Framework for Stable "
        "Modeling of Steady and Transient Navier–Stokes Flows.” Under review at IEEE Transactions "
        "on Neural Networks and Learning Systems (TNNLS).",
        size=10,
        after=2,
    )

    section_heading(doc, "Research & Engineering Experience")

    role_header(
        doc,
        "Head of Modern Swarm Warfare & Counter-UAS Systems",
        "Maxnet Labs (Defense / Aerospace R&D Initiatives) — Islamabad",
        "2023 – Present",
    )
    bullet(
        doc,
        "Anti-Drone System Architecture: Led design and integration of multi-sensor counter-UAS "
        "pipelines combining RF surveillance, direction finding, and signature analytics to detect, "
        "classify, and track unauthorized drones in contested electromagnetic environments.",
    )
    bullet(
        doc,
        "Anti-Drone Operator Detection: Developed RF-based operator localization and controller "
        "identification methods that fingerprint uplink/downlink emissions, isolate pilot-control "
        "links from ambient interference (Wi-Fi/Bluetooth), and support attribution of rogue "
        "operators beyond visual line of sight.",
    )
    bullet(
        doc,
        "Emitter Classification & Alerting: Applied machine-learning classifiers on transient and "
        "I/Q features to distinguish UAV platforms and controllers, improving detection robustness "
        "under low-SNR and multipath conditions typical of urban and industrial sites.",
    )
    bullet(
        doc,
        "Operational Concept Development: Authored CONOPS and technical proposals for layered "
        "counter-UAS response (detect → identify → geolocate → cue), including decentralized "
        "swarm-aware sensing concepts connecting RF fingerprinting with multi-agent coordination.",
    )

    role_header(
        doc,
        "RF & Sensor Systems Engineer",
        "Maxnet Labs (SUPARCO & Institute of Space Technology Collaboration)",
        "2022 – Present",
    )
    bullet(
        doc,
        "X-Band Sensor Development for Material Detection: Engineered high-frequency X-band "
        "microwave sensor systems for non-destructive material scanning and dielectric "
        "characterization, enabling penetration and contrast capabilities for structural and "
        "concealed-material analysis.",
    )
    bullet(
        doc,
        "Signal Processing & Attenuation Analysis: Processed complex backscatter and RF "
        "signatures to detect concealed materials; optimized penetration depth versus resolution "
        "trade-offs in high-noise environments.",
    )
    bullet(
        doc,
        "R&D Collaboration: Partnered with aerospace and applied-physics teams to integrate "
        "advanced sensor arrays into embedded operational frameworks and experimental testbeds.",
    )

    role_header(
        doc,
        "Principal Researcher — 6G Power Allocation & RF Fingerprinting",
        "NUST M.S. Thesis Research",
        "2025 – 2026",
    )
    bullet(
        doc,
        "Digital Twin Generation: Built a framework integrating drone-mounted LiDAR GSLAM to "
        "construct high-fidelity 3D digital twins of physical environments for spatially consistent "
        "RF simulation.",
    )
    bullet(
        doc,
        "Ray Tracing & CIR Computation: Used NVIDIA Sionna RT for differentiable ray tracing to "
        "compute Channel Impulse Responses (CIR) across dynamic spatial scenes.",
    )
    bullet(
        doc,
        "Machine Learning Integration: Coupled PINNs with Siamese CNNs to suppress multipath "
        "distortion and extract hardware signatures for 6G-related rogue-emitter identification "
        "and optimized power allocation.",
    )

    role_header(
        doc,
        "Researcher — Antennas & Wireless Systems",
        "Independent / IEEE-aligned Research Activities",
        "2016 – 2017",
    )
    bullet(
        doc,
        "Investigated WLAN optimization, smart antennas, 5G/terahertz antenna concepts, and "
        "material effects on wireless propagation as foundational preparation for later RF and "
        "sensing research.",
    )

    section_heading(doc, "Key Technical Projects")

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=1, line=1.0)
    r = p.add_run("Satellite Array Antenna Design & Fabrication")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    body_para(
        doc,
        "Led end-to-end design, simulation, and optimization of high-gain satellite antenna arrays "
        "in CST Microwave Studio—from CEM simulation and parameter tuning through physical "
        "prototyping and structural fabrication.",
        size=10,
        after=2,
    )

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=1, line=1.0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("Physics-Informed Neural Networks (PINNs) for Fluid Dynamics")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run("\tApr 2026")
    set_run_font(r2, size=9.5, color=MUTED)
    body_para(
        doc,
        "Architected scientific ML models using PINNs with SIREN activations and Time–Space "
        "Adaptive loss weighting. Processed JHU Turbulence Database parameters to mitigate "
        "spectral bias and model turbulent velocity fluctuations (basis of IEEE TNNLS submission).",
        size=10,
        after=2,
    )

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=1, line=1.0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("Hierarchical Post-Quantum Asynchronous Authentication")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run("\tApr 2026")
    set_run_font(r2, size=9.5, color=MUTED)
    body_para(
        doc,
        "Designed a zero-trust hardware integration framework for automotive enclaves. Drafted "
        "mathematical methodologies for Hardware Security Modules (HSMs) using lattice-based "
        "ML-KEM-768 master-seed encapsulation.",
        size=10,
        after=2,
    )

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=1, line=1.0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("Decentralized Drone Swarm Operations & Simulation")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run("\tAug 2026")
    set_run_font(r2, size=9.5, color=MUTED)
    body_para(
        doc,
        "Authored a research proposal mapping multi-agent systems with PINNs and Reinforcement "
        "Learning. Validated decentralized swarm environments using 3D multi-robot platforms "
        "(ARGoS, ROS, Gazebo), connecting swarm autonomy with RF-aware situational awareness.",
        size=10,
        after=2,
    )

    section_heading(doc, "Technical Skills")
    skills = [
        (
            "RF & Applied Electromagnetics: ",
            "X-band microwave sensors; NDT / dielectric characterization; satellite antenna arrays; "
            "6G architectures; RF fingerprinting; hardware-signature extraction; power allocation; "
            "counter-UAS RF detection; anti-drone operator localization.",
        ),
        (
            "Machine Learning & AI: ",
            "PINNs; Siamese CNNs; reinforcement learning; computer vision; scientific ML (SciML); "
            "emitter / UAV classification from RF features.",
        ),
        (
            "Simulation & Engineering Software: ",
            "CST Microwave Studio; MATLAB; NVIDIA Sionna RT; ROS; Gazebo; ARGoS; LiDAR / GSLAM "
            "processing; LaTeX.",
        ),
        (
            "Systems & Security: ",
            "Multi-agent / swarm simulation; post-quantum HSM concepts (ML-KEM-768); zero-trust "
            "hardware integration for automotive enclaves.",
        ),
    ]
    for label, content in skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=1, after=2, line=1.05)
        r1 = p.add_run(label)
        set_run_font(r1, size=10, bold=True, color=NAVY)
        r2 = p.add_run(content)
        set_run_font(r2, size=10, color=BODY)

    section_heading(doc, "Research Interests for Ph.D. Study")
    body_para(
        doc,
        "Applied electromagnetics and antenna systems; wireless channel modeling and measurement; "
        "RF / microwave sensing and computational imaging; RF fingerprinting and physical-layer "
        "security; scientific machine learning for wave and fluid physics; intelligent sensing for "
        "vehicular and aerospace wireless systems.",
        size=10,
        after=2,
    )


def build_cover_letter(doc):
    # Letterhead
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(name, before=0, after=0, line=1.0)
    r = name.add_run("USAMA MIR")
    set_run_font(r, size=16, bold=True, color=NAVY)

    contact = doc.add_paragraph()
    set_paragraph_spacing(contact, before=0, after=2, line=1.05)
    r = contact.add_run(
        "Islamabad, Pakistan  ·  payoneerusama@gmail.com  ·  linkedin.com/in/usama-mir-08678468"
    )
    set_run_font(r, size=9.5, color=MUTED)
    add_horizontal_line(contact, color="8B1E2D", thickness="12")

    date_p = doc.add_paragraph()
    set_paragraph_spacing(date_p, before=10, after=8, line=1.15)
    r = date_p.add_run("August 21, 2026")
    set_run_font(r, size=11, color=BODY)

    addr_lines = [
        "Graduate Admissions Committee",
        "Ph.D. Program in Electrical and Computer Engineering",
        "School of Engineering and Computer Science",
        "Oakland University",
        "Rochester, MI 48309",
        "United States",
    ]
    for i, line in enumerate(addr_lines):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=0 if i < len(addr_lines) - 1 else 8, line=1.1)
        r = p.add_run(line)
        set_run_font(r, size=11, color=BODY)

    salutation = doc.add_paragraph()
    set_paragraph_spacing(salutation, before=0, after=8, line=1.15)
    r = salutation.add_run("Dear Members of the Graduate Admissions Committee:")
    set_run_font(r, size=11, color=BODY)

    paragraphs = [
        (
            "I am writing to apply for admission to the Ph.D. program in Electrical and Computer "
            "Engineering at Oakland University. My academic and professional trajectory in RF "
            "systems, applied electromagnetics, counter-UAS sensing, and scientific machine learning "
            "has prepared me to contribute to Oakland’s research strengths in antennas, wireless "
            "propagation, and computational sensing—particularly within the Applied EMAG and "
            "Wireless Laboratory and related SECS initiatives in vehicular and aerospace wireless "
            "systems."
        ),
        (
            "My M.S. thesis at the National University of Sciences and Technology develops a hybrid "
            "framework that fuses LiDAR-based GSLAM digital twins, differentiable ray tracing with "
            "NVIDIA Sionna RT, Physics-Informed Neural Networks, and Siamese CNNs. The goal is "
            "practical and scientifically grounded: suppress multipath distortion in complex "
            "environments, recover hardware-level RF fingerprints, and support rogue-emitter "
            "identification and power-allocation decisions relevant to emerging 6G architectures. "
            "This work sits at the intersection of electromagnetics, channel modeling, and learning—"
            "precisely the kind of interdisciplinary problem space Oakland’s ECE doctoral program is "
            "designed to cultivate."
        ),
        (
            "Complementing this academic research, I have engineered X-band microwave sensors for "
            "non-destructive material characterization in collaboration with aerospace partners, and "
            "I have led counter-UAS and modern swarm-warfare efforts focused on anti-drone detection "
            "and anti-drone operator localization through RF surveillance and signature analytics. "
            "These experiences sharpened my ability to move from electromagnetic first principles to "
            "deployable sensing pipelines—processing noisy backscatter and control-link emissions, "
            "separating targets from interference, and translating algorithm design into operational "
            "concepts. Separately, my satellite array antenna work in CST Microwave Studio, from "
            "full-wave optimization through fabrication, reflects a sustained commitment to applied "
            "EM practice rather than simulation alone."
        ),
        (
            "I am equally invested in scientific machine learning as a tool for physical systems. My "
            "PINN research with SIREN activations and time–space adaptive loss weighting, trained "
            "against parameters from the Johns Hopkins Turbulence Database, addresses spectral bias "
            "in turbulent flow modeling and forms the basis of a manuscript under review at IEEE "
            "Transactions on Neural Networks and Learning Systems. I see strong synergy between this "
            "SciML agenda and RF problems involving inverse sensing, channel inference, and "
            "physics-constrained learning."
        ),
        (
            "Oakland University is my clear first-choice environment for doctoral study. The "
            "university’s depth in applied electromagnetics, antenna measurement and vehicular "
            "wireless testing, and industry-connected research culture offers the mentorship and "
            "facilities I need to scale my work on RF fingerprinting, computational sensors, and "
            "intelligent wireless systems. I would be honored to contribute as a doctoral researcher "
            "who bridges hardware, electromagnetics, and machine learning in service of safer and "
            "more capable sensing and communication technologies."
        ),
        (
            "Thank you for considering my application. I would welcome the opportunity to discuss "
            "how my background can support faculty research agendas in the School of Engineering "
            "and Computer Science. My résumé is enclosed for your review."
        ),
    ]

    for text in paragraphs:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=8, line=1.2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        set_run_font(r, size=11, color=BODY)

    close = doc.add_paragraph()
    set_paragraph_spacing(close, before=4, after=0, line=1.15)
    r = close.add_run("Sincerely,")
    set_run_font(r, size=11, color=BODY)

    # small photo + signature name
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.columns[0].width = Inches(1.3)
    sig_table.columns[1].width = Inches(5.5)
    c0 = sig_table.cell(0, 0)
    c1 = sig_table.cell(0, 1)
    c0.paragraphs[0].clear()
    if PHOTO.exists():
        run = c0.paragraphs[0].add_run()
        run.add_picture(str(PHOTO), width=Inches(0.95))
    c1.paragraphs[0].clear()
    p = c1.paragraphs[0]
    set_paragraph_spacing(p, before=18, after=0, line=1.1)
    r = p.add_run("Usama Mir")
    set_run_font(r, size=11, bold=True, color=NAVY)
    p2 = c1.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=0, line=1.1)
    r = p2.add_run("M.S. Electrical Engineering Candidate, NUST")
    set_run_font(r, size=10, color=MUTED)
    p3 = c1.add_paragraph()
    set_paragraph_spacing(p3, before=0, after=0, line=1.1)
    r = p3.add_run("Applicant, Ph.D. in Electrical & Computer Engineering")
    set_run_font(r, size=10, color=MUTED)
    p4 = c1.add_paragraph()
    set_paragraph_spacing(p4, before=0, after=0, line=1.1)
    r = p4.add_run("Oakland University")
    set_run_font(r, size=10, color=MUTED)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Usama Mir  ·  Oakland University Ph.D. Application  ·  Page ")
    set_run_font(run, size=8, color=MUTED)

    # PAGE field
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    set_run_font(run2, size=8, color=MUTED)
    run2._r.append(fldChar1)
    run2._r.append(instrText)
    run2._r.append(fldChar2)


def main():
    doc = Document()
    configure_page(doc.sections[0])

    # Cover letter
    build_cover_letter(doc)

    # Page break then resume
    doc.add_page_break()
    build_resume(doc)

    add_page_number(doc.sections[0])

    doc.save(OUTPUT)
    doc.save(ARTIFACT)
    print(f"Wrote {OUTPUT}")
    print(f"Wrote {ARTIFACT}")


if __name__ == "__main__":
    main()
